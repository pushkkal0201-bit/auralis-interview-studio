"""Bounded local extraction; uploads never become executable files."""
import io
import json
import zipfile
from pathlib import Path
from pypdf import PdfReader
from docx import Document

MAX_FILE=5*1024*1024

def extract_document(name,data):
    if len(data)>MAX_FILE:raise ValueError('Document exceeds 5 MB.')
    suffix=Path(name).suffix.lower()
    try:
        if suffix=='.txt':text=data.decode('utf-8-sig')
        elif suffix=='.pdf':
            if not data.startswith(b'%PDF'):raise ValueError('File is not a valid PDF.')
            reader=PdfReader(io.BytesIO(data))
            if reader.is_encrypted:raise ValueError('Unlock the PDF before uploading.')
            if len(reader.pages)>30:raise ValueError('Use a PDF with at most 30 pages.')
            parts=[]
            for page in reader.pages:
                parts.append(page.extract_text() or '')
                if sum(map(len,parts))>24000:raise ValueError('Document text exceeds 24,000 characters. Use a shorter CV.')
            text='\n'.join(parts)
        elif suffix=='.docx':
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if sum(v.file_size for v in z.infolist())>20*1024*1024:raise ValueError('Expanded DOCX is too large.')
            doc=Document(io.BytesIO(data));text='\n'.join(p.text for p in doc.paragraphs)
            text+='\n'+'\n'.join(' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows)
        else:raise ValueError('Upload a text-based PDF, DOCX or UTF-8 TXT file.')
    except ValueError:raise
    except Exception as e:raise ValueError('Cannot read this document. Try TXT or a text-based PDF.') from e
    text=text.strip()
    if not text:raise ValueError('No extractable text found. Paste the text from scanned documents.')
    if len(text)>24000:raise ValueError('Document text exceeds 24,000 characters. Use a shorter CV.')
    return text

def normalize_evidence(data):
    if len(data)>1024*1024:raise ValueError('REA JSON exceeds 1 MB.')
    try:obj=json.loads(data)
    except (ValueError,UnicodeError) as e:raise ValueError('Upload valid REA analysis JSON.') from e
    if not isinstance(obj,(dict,list)):raise ValueError('Analysis must be a JSON object or array.')
    text=json.dumps(obj,ensure_ascii=False,indent=2)
    if len(text)>12000:raise ValueError('Select a smaller analysis excerpt, at most 12,000 characters.')
    return text
